from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = ROOT / "analysis" / "outputs" / "2.2 Quadro da Teoria da Mudanca"
SOURCE = OUT_DIR / "TABELA 31 - Teoria da Mudanca - 2023 a 2026.xlsx"
OUT = OUT_DIR / "TABELA 31 - Teoria da Mudanca - analise por cnpj completo.xlsx"


def style(path: Path) -> None:
    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="7F7F7F")
    header_font = Font(color="FFFFFF", bold=True)
    stripe_fill = PatternFill("solid", fgColor="F4F7F6")
    side = Side(style="thin", color="B7B7B7")
    border = Border(left=side, right=side, top=side, bottom=side)
    for ws in wb.worksheets:
        if ws.max_row < 1 or ws.max_column < 1:
            continue
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ws.max_column):
            if row[0].row % 2 == 0:
                for cell in row:
                    cell.fill = stripe_fill
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for idx in range(1, ws.max_column + 1):
            letter = get_column_letter(idx)
            max_len = 10
            for cell in ws[letter][: min(ws.max_row, 100)]:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 60)
    wb.save(path)


def main() -> None:
    panel = pd.read_excel(SOURCE, sheet_name="painel_empresas_rais", dtype={"cnpj_completo": str})
    treated = panel[panel["papel"].eq("tratada")].copy()
    id_cols = [
        "cnpj_completo",
        "cnpj_raiz",
        "cnpj_key",
        "cliente",
        "ano_primeira_operacao",
        "tipo_operacao",
        "regiao",
        "setor",
    ]
    base = treated.sort_values(["cnpj_completo", "ano"]).groupby("cnpj_completo", dropna=False).agg(
        {col: "first" for col in id_cols if col != "cnpj_completo"}
    ).reset_index()
    estoque = treated.pivot_table(
        index="cnpj_completo",
        columns="ano",
        values="estoque_medio_anual_empregados",
        aggfunc="first",
    ).add_prefix("estoque_medio_").reset_index()
    variacao = treated.pivot_table(
        index="cnpj_completo",
        columns="ano",
        values="variacao_pct_empregados",
        aggfunc="first",
    ).add_prefix("variacao_").reset_index()
    resumo = base.merge(estoque, on="cnpj_completo", how="left").merge(
        variacao, on="cnpj_completo", how="left"
    )
    with pd.ExcelWriter(OUT, engine="openpyxl") as writer:
        resumo.to_excel(writer, index=False, sheet_name="resumo_cnpj")
        treated.to_excel(writer, index=False, sheet_name="painel_cnpj_ano")
        pd.DataFrame(
            [
                {
                    "nota": "Analise no nivel de CNPJ completo. A variacao anual usa o estoque medio mensal anual do proprio CNPJ completo."
                }
            ]
        ).to_excel(writer, index=False, sheet_name="metodologia")
    style(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
