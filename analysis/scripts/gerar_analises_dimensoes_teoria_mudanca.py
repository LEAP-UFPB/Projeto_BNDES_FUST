from __future__ import annotations

import re
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = ROOT / "analysis" / "outputs" / "2.2 Quadro da Teoria da Mudanca"
SOURCE = OUT_DIR / "TABELA 31 - Teoria da Mudanca - 2023 a 2026.xlsx"

DIMENSIONS = {
    "Regiao": "regiao",
    "Porte": "porte",
    "Setor": "setor",
}

LABELS = {
    ("I1", "valor observado", "nao aplicavel"): "I1 - numero de operacoes aprovadas",
    ("I2", "valor observado", "nao aplicavel"): "I2 - numero de clientes apoiados",
    ("I3", "valor observado", "nao aplicavel"): "I3 - novos clientes",
    ("I6", "mediana", "ano de referencia vs anterior"): "I6.med.ant - mediana da variacao percentual do estoque medio anual de empregados das empresas apoiadas - ano de referencia vs anterior",
    ("I6", "mediana", "ano posterior vs referencia"): "I6.med.pos - mediana da variacao percentual do estoque medio anual de empregados das empresas apoiadas - ano posterior vs referencia",
    ("I6", "media", "ano de referencia vs anterior"): "I6.if5.ant - media da variacao percentual do estoque medio anual de empregados das empresas apoiadas - ano de referencia vs anterior",
    ("I6", "media", "ano posterior vs referencia"): "I6.if5.pos - media da variacao percentual do estoque medio anual de empregados das empresas apoiadas - ano posterior vs referencia",
    ("I7", "mediana", "ano de referencia vs anterior"): "I7.med.ant - mediana da variacao percentual do estoque medio anual de empregados das empresas nao apoiadas pareadas - ano de referencia vs anterior",
    ("I7", "mediana", "ano posterior vs referencia"): "I7.med.pos - mediana da variacao percentual do estoque medio anual de empregados das empresas nao apoiadas pareadas - ano posterior vs referencia",
    ("I7", "media", "ano de referencia vs anterior"): "I7.ant - media da variacao percentual do estoque medio anual de empregados das empresas nao apoiadas pareadas - ano de referencia vs anterior",
    ("I7", "media", "ano posterior vs referencia"): "I7.pos - media da variacao percentual do estoque medio anual de empregados das empresas nao apoiadas pareadas - ano posterior vs referencia",
}


def fmt_pct(value: object) -> str:
    if pd.isna(value):
        return "*"
    return f"{float(value) * 100:.1f}%".replace(".", ",")


def fmt_number(value: object) -> str:
    if pd.isna(value):
        return "*"
    return f"{int(round(float(value))):,}".replace(",", ".")


def fmt_value(codigo: str, value: object) -> str:
    if codigo in {"I6", "I7"}:
        return fmt_pct(value)
    return fmt_number(value)


def build_summary(rows: pd.DataFrame, tipo: str | None) -> pd.DataFrame:
    base = rows.copy() if tipo is None else rows[rows["tipo_operacao"].eq(tipo)].copy()
    years = sorted(rows["ano_aprovacao"].dropna().astype(int).unique())
    records = []
    categories = sorted(base["categoria"].dropna().astype(str).unique())
    for category in categories:
        for codigo, stat, janela in [
            ("I1", "valor observado", "nao aplicavel"),
            ("I2", "valor observado", "nao aplicavel"),
            ("I3", "valor observado", "nao aplicavel"),
            ("I6", "mediana", "ano de referencia vs anterior"),
            ("I6", "mediana", "ano posterior vs referencia"),
            ("I6", "media", "ano de referencia vs anterior"),
            ("I6", "media", "ano posterior vs referencia"),
            ("I7", "mediana", "ano de referencia vs anterior"),
            ("I7", "mediana", "ano posterior vs referencia"),
            ("I7", "media", "ano de referencia vs anterior"),
            ("I7", "media", "ano posterior vs referencia"),
        ]:
            subset = base[
                base["codigo_indicador"].eq(codigo)
                & base["estatistica"].eq(stat)
                & base["janela_calculo"].eq(janela)
                & base["categoria"].astype(str).eq(category)
            ].copy()
            rec = {
                "Categoria": category,
                "Indicador": LABELS[(codigo, stat, janela)],
            }
            for year in years:
                yr = subset[subset["ano_aprovacao"].astype("Int64").eq(year)]
                rec[str(year)] = "*" if yr.empty else fmt_value(codigo, yr.iloc[0]["valor_indicador"])
            records.append(rec)
    return pd.DataFrame(records)


def safe_sheet_name(value: str) -> str:
    value = re.sub(r"[\[\]\:\*\?\/\\]", " ", value)
    value = re.sub(r"\s+", " ", value).strip()
    return value[:31] or "Planilha"


def style(path: Path) -> None:
    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="7F7F7F")
    header_font = Font(color="FFFFFF", bold=True)
    stripe_fill = PatternFill("solid", fgColor="F4F7F6")
    side = Side(style="thin", color="B7B7B7")
    border = Border(left=side, right=side, top=side, bottom=side)
    for ws in wb.worksheets:
        if ws.max_row < 1 or ws.max_column < 1:
            continue
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ws.max_column):
            if row[0].row % 2 == 0:
                for cell in row:
                    cell.fill = stripe_fill
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for idx in range(1, ws.max_column + 1):
            letter = get_column_letter(idx)
            max_len = 10
            for cell in ws[letter][: min(ws.max_row, 100)]:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 65)
    wb.save(path)


def write_dimension(dimension: str, suffix: str, por_ano: pd.DataFrame, audit: pd.DataFrame, panel: pd.DataFrame) -> Path:
    rows = por_ano[
        por_ano["dimensao"].eq(dimension)
        & por_ano["codigo_indicador"].isin(["I1", "I2", "I3", "I6", "I7"])
    ].copy()
    audit_rows = audit[audit["dimensao"].eq(dimension)].copy()

    out = OUT_DIR / f"TABELA 31 - Teoria da Mudanca - analise por {suffix}.xlsx"
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        build_summary(rows, "Total").to_excel(writer, index=False, sheet_name="Total")
        build_summary(rows, "Direta").to_excel(writer, index=False, sheet_name="Direta")
        build_summary(rows, "Indireta").to_excel(writer, index=False, sheet_name="Indireta")
        rows.sort_values(
            ["categoria", "tipo_operacao", "ano_aprovacao", "codigo_indicador", "estatistica"]
        ).to_excel(writer, index=False, sheet_name="dados")
        audit_rows.sort_values(
            ["categoria", "tipo_operacao", "ano_aprovacao"]
        ).to_excel(writer, index=False, sheet_name="auditoria_matching")
        panel.to_excel(writer, index=False, sheet_name="painel_empresas_rais")
        pd.DataFrame(
            [
                {
                    "nota": (
                        f"Analise por {suffix}. I1/I2/I3 usam ano de aprovacao quando disponivel. "
                        "I6/I7 usam o painel completo das empresas apoiadas e controles com estoque medio mensal anual RAIS. "
                        "Anos sem RAIS disponivel ficam como *."
                    )
                }
            ]
        ).to_excel(writer, index=False, sheet_name="metodologia")
    style(out)
    return out


def main() -> None:
    por_ano = pd.read_excel(SOURCE, sheet_name="por_ano")
    audit = pd.read_excel(SOURCE, sheet_name="auditoria_matching")
    panel = pd.read_excel(SOURCE, sheet_name="painel_empresas_rais")
    for dimension, suffix in DIMENSIONS.items():
        path = write_dimension(dimension, suffix, por_ano, audit, panel)
        print(path)


if __name__ == "__main__":
    main()
