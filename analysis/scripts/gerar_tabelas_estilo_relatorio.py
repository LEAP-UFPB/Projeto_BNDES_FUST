from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = ROOT / "analysis" / "outputs" / "2.2 Quadro da Teoria da Mudanca"
SOURCE = OUT_DIR / (
    "TABELA 31 - Teoria da Mudanca - 2023 a 2026.xlsx"
)
OUT = OUT_DIR / (
    "TABELA 31 - Teoria da Mudanca - estilo relatorio.xlsx"
)


INDICATOR_LABEL = {
    "I1": "I1 - numero de operacoes aprovadas",
    "I2": "I2 - numero de clientes apoiados",
    "I3": "I3 - novos clientes",
    "I4": "I4 - equipamentos adquiridos ou modernizados por tipo",
    "I5": "I5 - area construida ou modernizada em projetos de investimento",
    "I6": "I6 - variacao percentual do estoque medio anual de empregados das empresas apoiadas",
    "I7": "I7 - indicador de controle para comparacao",
    "I8": "I8 - variacao percentual do faturamento",
    "I9": "I9 - variacao percentual do investimento fixo das empresas apoiadas",
}


def fmt_number(value: object) -> str:
    if pd.isna(value):
        return "*"
    value = float(value)
    if abs(value - round(value)) < 1e-9:
        return f"{int(round(value)):,}".replace(",", ".")
    return f"{value:,.1f}".replace(",", "X").replace(".", ",").replace("X", ".")


def fmt_pct(value: object) -> str:
    if pd.isna(value):
        return "*"
    return f"{float(value) * 100:.1f}%".replace(".", ",")


def fmt_cell(row: pd.Series, value_col: str = "valor_indicador") -> str:
    value = row[value_col]
    codigo = row["codigo_indicador"]
    if pd.isna(value):
        return "*"
    if codigo in {"I6", "I7", "I8", "I9"}:
        return fmt_pct(value)
    return fmt_number(value)


def prepare_rows(por_ano: pd.DataFrame, tipo: str | None) -> pd.DataFrame:
    if tipo is None:
        base = por_ano.copy()
    else:
        base = por_ano[por_ano["tipo_operacao"].eq(tipo)].copy()

    rows = []
    years = sorted(por_ano["ano_aprovacao"].dropna().astype(int).unique())

    def add_row(
        result: str,
        codigo: str,
        label: str,
        stat: str,
        dim: str = "Geral",
        cat: str = "Total",
        janela: str | None = None,
    ):
        subset = base[
            base["codigo_indicador"].eq(codigo)
            & base["estatistica"].eq(stat)
            & base["dimensao"].eq(dim)
            & base["categoria"].eq(cat)
        ].copy()
        if janela is not None and "janela_calculo" in subset.columns:
            subset = subset[subset["janela_calculo"].eq(janela)].copy()
        if tipo is None and not subset.empty:
            # Total: for counts, sum direct+indirect; for percentage indicators, average
            # across operation types weighted implicitly by available rows in this report table.
            grouped_rows = []
            for year, g in subset.groupby("ano_aprovacao"):
                if codigo in {"I1", "I2", "I3"}:
                    val = pd.to_numeric(g["valor_indicador"], errors="coerce").sum(min_count=1)
                    n = pd.to_numeric(g["clientes"], errors="coerce").sum(min_count=1)
                elif codigo in {"I6", "I7"}:
                    val = pd.to_numeric(g["valor_indicador"], errors="coerce").mean()
                    n = pd.to_numeric(
                        g["clientes"] if codigo == "I6" else g["controles_distintos"],
                        errors="coerce",
                    ).sum(min_count=1)
                else:
                    val = pd.NA
                    n = pd.NA
                grouped_rows.append({"ano_aprovacao": year, "valor_indicador": val, "n_linha": n})
            subset = pd.DataFrame(grouped_rows)
        elif not subset.empty:
            subset["n_linha"] = subset.apply(
                lambda r: r["clientes"] if codigo in {"I1", "I2", "I3", "I6"} else (
                    r["controles_distintos"] if codigo == "I7" else pd.NA
                ),
                axis=1,
            )
        else:
            subset = pd.DataFrame(columns=["ano_aprovacao", "valor_indicador", "n_linha"])

        rec = {"Resultado": result, "Indicador": label}
        for year in years:
            yr = subset[subset["ano_aprovacao"].eq(year)]
            if yr.empty:
                rec[str(year)] = "*"
            else:
                row = yr.iloc[0].copy()
                row["codigo_indicador"] = codigo
                rec[str(year)] = fmt_cell(row)
        rows.append(rec)

    add_row("Eficacia", "I1", INDICATOR_LABEL["I1"], "valor observado")
    add_row("Eficacia", "I2", INDICATOR_LABEL["I2"], "valor observado")
    add_row("Eficacia", "I3", INDICATOR_LABEL["I3"], "valor observado")
    add_row("Eficacia", "I4", INDICATOR_LABEL["I4"], "nao disponivel")
    add_row("Eficacia", "I5", INDICATOR_LABEL["I5"], "nao disponivel")
    add_row(
        "Efetividade",
        "I6",
        "I6.med.ant - mediana da variacao percentual do estoque medio anual de empregados das empresas apoiadas - ano de referencia vs anterior",
        "mediana",
        janela="ano de referencia vs anterior",
    )
    add_row(
        "Efetividade",
        "I6",
        "I6.med.pos - mediana da variacao percentual do estoque medio anual de empregados das empresas apoiadas - ano posterior vs referencia",
        "mediana",
        janela="ano posterior vs referencia",
    )
    add_row(
        "Efetividade",
        "I6",
        "I6.if5.ant - media da variacao percentual do estoque medio anual de empregados das empresas apoiadas - ano de referencia vs anterior",
        "media",
        janela="ano de referencia vs anterior",
    )
    add_row(
        "Efetividade",
        "I6",
        "I6.if5.pos - media da variacao percentual do estoque medio anual de empregados das empresas apoiadas - ano posterior vs referencia",
        "media",
        janela="ano posterior vs referencia",
    )
    add_row(
        "Efetividade",
        "I7",
        "I7.med.ant - mediana da variacao percentual do estoque medio anual de empregados das empresas nao apoiadas pareadas - ano de referencia vs anterior",
        "mediana",
        janela="ano de referencia vs anterior",
    )
    add_row(
        "Efetividade",
        "I7",
        "I7.med.pos - mediana da variacao percentual do estoque medio anual de empregados das empresas nao apoiadas pareadas - ano posterior vs referencia",
        "mediana",
        janela="ano posterior vs referencia",
    )
    add_row(
        "Efetividade",
        "I7",
        "I7.ant - media da variacao percentual do estoque medio anual de empregados das empresas nao apoiadas pareadas - ano de referencia vs anterior",
        "media",
        janela="ano de referencia vs anterior",
    )
    add_row(
        "Efetividade",
        "I7",
        "I7.pos - media da variacao percentual do estoque medio anual de empregados das empresas nao apoiadas pareadas - ano posterior vs referencia",
        "media",
        janela="ano posterior vs referencia",
    )
    add_row("Efetividade", "I8", INDICATOR_LABEL["I8"], "nao disponivel")
    add_row("Efetividade", "I9", INDICATOR_LABEL["I9"], "nao disponivel")
    return pd.DataFrame(rows)


def indicator_sheet(por_ano: pd.DataFrame, codigo: str) -> pd.DataFrame:
    years = sorted(por_ano["ano_aprovacao"].dropna().astype(int).unique())
    parts = []
    for tipo in ["Direta", "Indireta", "Total"]:
        table = prepare_rows(por_ano, tipo)
        table = table[table["Indicador"].str.startswith(f"{codigo}")]
        table.insert(0, "Tipo", tipo)
        parts.append(table)
    if parts:
        return pd.concat(parts, ignore_index=True)
    return pd.DataFrame(columns=["Tipo", "Resultado", "Indicador", *map(str, years)])


def style(path: Path) -> None:
    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="7F7F7F")
    header_font = Font(color="FFFFFF", bold=True)
    result_fill = PatternFill("solid", fgColor="EAF2F1")
    side = Side(style="thin", color="000000")
    border = Border(left=side, right=side, top=side, bottom=side)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ws.max_column):
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            if row[0].value in {"Eficacia", "Efetividade"} or (
                ws.max_column > 1 and row[1].value in {"Eficacia", "Efetividade"}
            ):
                row[0].fill = result_fill
        widths = {"A": 16, "B": 85}
        if ws.cell(row=1, column=1).value == "Tipo":
            widths = {"A": 14, "B": 16, "C": 85}
        for idx in range(1, ws.max_column + 1):
            letter = get_column_letter(idx)
            ws.column_dimensions[letter].width = widths.get(letter, 18)
        for row_idx in range(1, ws.max_row + 1):
            ws.row_dimensions[row_idx].height = 34 if row_idx > 1 else 24
    wb.save(path)


def main() -> None:
    por_ano = pd.read_excel(SOURCE, sheet_name="por_ano")
    out = OUT
    try:
        with open(out, "a+b"):
            pass
    except PermissionError:
        out = OUT.with_name(OUT.stem + " - atualizado.xlsx")

    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        tables = {
            "Direta": prepare_rows(por_ano, "Direta"),
            "Indireta": prepare_rows(por_ano, "Indireta"),
            "Total": prepare_rows(por_ano, "Total"),
        }
        for name, df in tables.items():
            df.to_excel(writer, index=False, sheet_name=name)
        for codigo in [f"I{i}" for i in range(1, 10)]:
            indicator_sheet(por_ano, codigo).to_excel(
                writer,
                index=False,
                sheet_name=codigo,
            )
    style(out)
    print(out)


if __name__ == "__main__":
    main()
