from __future__ import annotations

import math
import re
import unicodedata
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[2]
DATA_DIR = ROOT / "data"
OUTPUT_DIR = ROOT / "analysis" / "outputs" / "2.2 Quadro da Teoria da Mudanca"

TITULO_BASE = "TABELA 31 - Teoria da Mudanca"

UF_REGIAO = {
    "AC": "Norte",
    "AP": "Norte",
    "AM": "Norte",
    "PA": "Norte",
    "RO": "Norte",
    "RR": "Norte",
    "TO": "Norte",
    "AL": "Nordeste",
    "BA": "Nordeste",
    "CE": "Nordeste",
    "MA": "Nordeste",
    "PB": "Nordeste",
    "PE": "Nordeste",
    "PI": "Nordeste",
    "RN": "Nordeste",
    "SE": "Nordeste",
    "DF": "Centro-Oeste",
    "GO": "Centro-Oeste",
    "MS": "Centro-Oeste",
    "MT": "Centro-Oeste",
    "ES": "Sudeste",
    "MG": "Sudeste",
    "RJ": "Sudeste",
    "SP": "Sudeste",
    "PR": "Sul",
    "RS": "Sul",
    "SC": "Sul",
}


def only_digits(value: object) -> str:
    if pd.isna(value):
        return ""
    return re.sub(r"\D", "", str(value))


def mask_cnpj_from_digits(value: object) -> str:
    digits = only_digits(value)
    if len(digits) != 14:
        return ""
    return f"**.*{digits[3:5]}.{digits[5:8]}/{digits[8:12]}-**"


def normalize_cnpj_piece(value: object) -> str:
    if pd.isna(value):
        return ""
    return re.sub(r"\s+", "", str(value).strip()).upper()


def cnpj_key_set(value: object) -> set[str]:
    raw = "" if pd.isna(value) else str(value).strip().upper()
    parts = [part.strip() for part in raw.split("|") if part.strip()]
    keys: set[str] = set()
    for part in parts or [raw]:
        compact = normalize_cnpj_piece(part)
        digits = only_digits(compact)
        if "*" in compact:
            keys.add(compact)
        if len(digits) == 14:
            keys.add(mask_cnpj_from_digits(digits))
    return {key for key in keys if key}


def primary_cnpj_key(value: object) -> str:
    keys = cnpj_key_set(value)
    masked = sorted(key for key in keys if "*" in key)
    if masked:
        return masked[0]
    return sorted(keys)[0] if keys else ""


def normalize_text(value: object) -> str:
    if pd.isna(value):
        return ""
    text = str(value).upper().strip()
    text = "".join(
        char
        for char in unicodedata.normalize("NFKD", text)
        if not unicodedata.combining(char)
    )
    text = re.sub(r"[^A-Z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def first_non_empty(values: pd.Series) -> object:
    values = values.dropna()
    values = values[values.astype(str).str.strip().ne("")]
    return values.iloc[0] if len(values) else pd.NA


def load_approval_dates_from_qr() -> pd.DataFrame:
    parts = []
    for path in sorted(DATA_DIR.glob("Base de QR FUST *.xlsx")):
        try:
            qr = pd.read_excel(path, sheet_name=0)
        except Exception:
            continue
        required = {"operação número", "data.aprovacao"}
        if not required.issubset(set(qr.columns)):
            continue
        part = qr[["operação número", "data.aprovacao"]].copy()
        part["numero_de_operacao_key"] = pd.to_numeric(
            part["operação número"], errors="coerce"
        ).astype("Int64")
        part["data_aprovacao_qr"] = pd.to_datetime(
            part["data.aprovacao"], errors="coerce"
        )
        part = part[
            part["numero_de_operacao_key"].notna()
            & part["data_aprovacao_qr"].notna()
        ][["numero_de_operacao_key", "data_aprovacao_qr"]]
        parts.append(part)
    if not parts:
        return pd.DataFrame(columns=["numero_de_operacao_key", "data_aprovacao_qr"])
    approval = pd.concat(parts, ignore_index=True).drop_duplicates()
    return (
        approval.groupby("numero_de_operacao_key", dropna=False)["data_aprovacao_qr"]
        .min()
        .reset_index()
    )


def employment_band(value: object) -> str:
    if pd.isna(value):
        return "Sem informacao"
    value = float(value)
    if value <= 0:
        return "0"
    if value <= 9:
        return "1 a 9"
    if value <= 49:
        return "10 a 49"
    if value <= 99:
        return "50 a 99"
    if value <= 249:
        return "100 a 249"
    if value <= 499:
        return "250 a 499"
    return "500 ou mais"


def safe_mean(values: list[float]) -> float:
    clean = [value for value in values if value is not None and not pd.isna(value)]
    return sum(clean) / len(clean) if clean else math.nan


def safe_median(values: list[float]) -> float:
    clean = sorted(value for value in values if value is not None and not pd.isna(value))
    if not clean:
        return math.nan
    mid = len(clean) // 2
    if len(clean) % 2:
        return clean[mid]
    return (clean[mid - 1] + clean[mid]) / 2


def safe_quantile(values: list[float], q: float) -> float:
    clean = sorted(value for value in values if value is not None and not pd.isna(value))
    if not clean:
        return math.nan
    if len(clean) == 1:
        return clean[0]
    pos = (len(clean) - 1) * q
    low = math.floor(pos)
    high = math.ceil(pos)
    if low == high:
        return clean[int(pos)]
    return clean[low] + (clean[high] - clean[low]) * (pos - low)


def parse_base_operacoes() -> pd.DataFrame:
    base_path = next(DATA_DIR.glob("Base de Opera*compilada.xlsx"))
    base = pd.read_excel(base_path, sheet_name="base_compilada")
    base = base.reset_index(names="linha_base")
    base["data_contratacao"] = pd.to_datetime(
        base["data_da_contratacao"], dayfirst=True, errors="coerce"
    )
    base["numero_de_operacao_key"] = pd.to_numeric(
        base["numero_de_operacao"], errors="coerce"
    ).astype("Int64")
    approval = load_approval_dates_from_qr()
    base = base.merge(approval, on="numero_de_operacao_key", how="left")
    base["data_aprovacao"] = base["data_aprovacao_qr"]
    base["fonte_data_aprovacao"] = "QR FUST"
    missing_approval = base["data_aprovacao"].isna()
    base.loc[missing_approval, "data_aprovacao"] = base.loc[
        missing_approval, "data_contratacao"
    ]
    base.loc[missing_approval, "fonte_data_aprovacao"] = (
        "fallback: data_da_contratacao sem data.aprovacao no QR"
    )
    base["ano_aprovacao"] = base["data_aprovacao"].dt.year.astype("Int64")
    base["cnpj_key"] = base["cnpj"].map(primary_cnpj_key)
    base["cnpj_digits"] = base["cnpj"].map(only_digits)
    base["cnpj_completo_base"] = base["cnpj_digits"].where(
        base["cnpj_digits"].str.len().eq(14),
        pd.NA,
    )
    base["cnpj_raiz_base"] = base["cnpj_digits"].where(
        base["cnpj_digits"].str.len().eq(14),
        pd.NA,
    ).str[:8]
    base["cliente_norm"] = base["cliente"].map(normalize_text)
    base["tipo_operacao"] = base["tipo_operacao"].fillna("Nao informado")
    base["uf"] = base["uf"].fillna("Nao informado").astype(str).str.upper()
    base["regiao"] = base["uf"].map(UF_REGIAO).fillna("Nao informado")

    contrato_num = pd.to_numeric(base["numero_do_contrato"], errors="coerce")
    operacao_num = pd.to_numeric(base["numero_de_operacao"], errors="coerce")
    base["contrato_id"] = ""
    has_contrato = contrato_num.notna()
    has_operacao = ~has_contrato & operacao_num.notna()
    base.loc[has_contrato, "contrato_id"] = (
        "C_" + contrato_num[has_contrato].astype("int64").astype(str)
    )
    base.loc[has_operacao, "contrato_id"] = (
        "O_" + operacao_num[has_operacao].astype("int64").astype(str)
    )
    base.loc[base["contrato_id"].eq(""), "contrato_id"] = (
        "L_" + base.loc[base["contrato_id"].eq(""), "linha_base"].astype(str)
    )

    return base


def company_profile(base: pd.DataFrame) -> pd.DataFrame:
    ordered = base.sort_values(["cnpj_key", "data_aprovacao", "linha_base"])
    profile = (
        ordered.groupby("cnpj_key", dropna=False)
        .agg(
            cliente=("cliente", first_non_empty),
            cnpj_completo_base=("cnpj_completo_base", first_non_empty),
            cnpj_raiz_base=("cnpj_raiz_base", first_non_empty),
            ano_primeira_operacao=("ano_aprovacao", "min"),
            data_primeira_operacao=("data_aprovacao", "min"),
            tipo_primeira_operacao=("tipo_operacao", first_non_empty),
            uf_fust=("uf", first_non_empty),
            regiao_fust=("regiao", first_non_empty),
            porte_fust=("porte_do_cliente", first_non_empty),
            setor_cnae_fust=("setor_cnae", first_non_empty),
            subsetor_cnae_fust=("subsetor_cnae_agrupado", first_non_empty),
            valor_operacao_fust=("valor_operacao_r", "sum"),
            valor_desembolsado_fust=("valor_desembolsado_r", "sum"),
            qtd_contratos_fust=("contrato_id", "nunique"),
        )
        .reset_index()
    )
    return profile


def contract_view(base: pd.DataFrame) -> pd.DataFrame:
    contracts = (
        base.sort_values(["contrato_id", "ano_aprovacao", "data_aprovacao", "linha_base"])
        .groupby(["contrato_id", "ano_aprovacao"], dropna=False)
        .agg(
            data_aprovacao=("data_aprovacao", "min"),
            data_contratacao=("data_contratacao", "min"),
            fonte_data_aprovacao=("fonte_data_aprovacao", first_non_empty),
            tipo_operacao=("tipo_operacao", first_non_empty),
            cnpj_key=("cnpj_key", first_non_empty),
            cliente=("cliente", first_non_empty),
            uf=("uf", first_non_empty),
            regiao=("regiao", first_non_empty),
            porte=("porte_do_cliente", first_non_empty),
            setor=("subsetor_cnae_agrupado", first_non_empty),
            valor_operacao_r=("valor_operacao_r", "sum"),
            valor_desembolsado_r=("valor_desembolsado_r", "sum"),
            linhas_base=("linha_base", "size"),
        )
        .reset_index()
    )
    contracts["ano_aprovacao"] = contracts["ano_aprovacao"].astype("Int64")
    return contracts


def parse_rais(profile: pd.DataFrame) -> pd.DataFrame:
    painel_path = DATA_DIR / "Painel_CNPJ_FUST_RAIS_2018_2024.xlsx"
    rais = pd.read_excel(painel_path, sheet_name="painel_cnpj_ano")
    rais["cnpj_key"] = rais["cnpj_fust"].map(primary_cnpj_key)
    rais = rais[rais["cnpj_key"].ne("")].copy()
    rais = rais.sort_values(["cnpj_key", "ano"])
    rais = rais.merge(
        profile[
            [
                "cnpj_key",
                "ano_primeira_operacao",
                "tipo_primeira_operacao",
                "regiao_fust",
                "porte_fust",
                "setor_cnae_fust",
                "subsetor_cnae_fust",
            ]
        ],
        on="cnpj_key",
        how="left",
        suffixes=("", "_base"),
    )

    rais["total_empregados"] = pd.to_numeric(rais["total_empregados"], errors="coerce")
    rais["empregados_lag"] = rais.groupby("cnpj_key")["total_empregados"].shift(1)
    rais["ano_lag"] = rais.groupby("cnpj_key")["ano"].shift(1)
    rais["variacao_pct_empregados"] = (
        rais["total_empregados"] - rais["empregados_lag"]
    ) / rais["empregados_lag"]
    rais.loc[
        (rais["empregados_lag"].isna())
        | (rais["empregados_lag"] <= 0)
        | (rais["ano_lag"] != rais["ano"] - 1),
        "variacao_pct_empregados",
    ] = pd.NA
    rais["faixa_emprego_pre"] = rais["empregados_lag"].map(employment_band)
    rais["regiao"] = rais["regiao_fust"].fillna(
        rais["uf_fust"].astype(str).str.upper().map(UF_REGIAO)
    )
    rais["porte"] = rais["porte_fust"].fillna("Nao informado")
    rais["setor"] = rais["setor_cnae_fust"].fillna("Nao informado")
    rais["subsetor"] = rais["subsetor_cnae_fust"].fillna("Nao informado")
    return rais


def region_from_uf(uf: object) -> str:
    if pd.isna(uf):
        return "Nao informado"
    return UF_REGIAO.get(str(uf).upper(), "Nao informado")


def cnpj_key_to_root_map() -> pd.DataFrame:
    painel_path = DATA_DIR / "Painel_CNPJ_FUST_RAIS_2018_2024.xlsx"
    painel = pd.read_excel(painel_path, sheet_name="painel_cnpj_ano")
    rows = []
    for _, row in painel[["cnpj_fust", "cnpj_raiz"]].drop_duplicates().iterrows():
        if pd.isna(row["cnpj_raiz"]):
            continue
        root = str(int(row["cnpj_raiz"])).zfill(8)
        for key in cnpj_key_set(row["cnpj_fust"]):
            if "*" in key:
                rows.append({"cnpj_key": key, "cnpj_raiz": root})
    return pd.DataFrame(rows).drop_duplicates()


def cnpj_key_to_full_map() -> pd.DataFrame:
    painel_path = DATA_DIR / "Painel_CNPJ_FUST_RAIS_2018_2024.xlsx"
    painel = pd.read_excel(painel_path, sheet_name="painel_cnpj_ano")
    rows = []
    cols = ["cnpj_fust", "cnpj_matriz_rais"]
    for _, row in painel[cols].drop_duplicates().iterrows():
        full = only_digits(row["cnpj_matriz_rais"])
        if len(full) != 14:
            continue
        for key in cnpj_key_set(row["cnpj_fust"]):
            if "*" in key:
                rows.append({"cnpj_key": key, "cnpj_completo": full})
    return pd.DataFrame(rows).drop_duplicates()


def parse_rais_reprocessada(
    profile: pd.DataFrame,
    years: list[int],
) -> tuple[pd.DataFrame, pd.DataFrame, dict[str, object]]:
    universo_path = DATA_DIR / "RAIS_universo_empresas_2018_2024.csv"
    if not universo_path.exists():
        raise FileNotFoundError(
            f"{universo_path} nao encontrado. Rode reprocess_rais_universo_empresas.R antes."
        )

    key_root = cnpj_key_to_root_map()
    profile_root = profile.merge(key_root, on="cnpj_key", how="left")
    profile_root["cnpj_raiz"] = profile_root["cnpj_raiz"].fillna(profile_root["cnpj_raiz_base"])
    missing_roots = int(profile_root["cnpj_raiz"].isna().sum())
    profile_root = profile_root[profile_root["cnpj_raiz"].notna()].copy()
    profile_root["cnpj_raiz"] = profile_root["cnpj_raiz"].astype(str).str.zfill(8)
    fust_roots = set(profile_root["cnpj_raiz"])

    use_years = [year for year in years if year <= 2024]
    usecols = [
        "ano",
        "cnpj_raiz",
        "total_empregados",
        "empregados_lag",
        "variacao_pct_empregados",
        "faixa_emprego_pre",
        "uf_principal",
        "setor_cnae_classe_principal",
        "setor_cnae_subclasse_principal",
        "subsetor_ibge_principal_codigo",
    ]

    treated_parts = []
    control_parts = []
    total_rows = 0
    companies_seen: set[str] = set()
    chunks = pd.read_csv(
        universo_path,
        sep=";",
        usecols=usecols,
        dtype={
            "ano": "int16",
            "cnpj_raiz": "string",
            "faixa_emprego_pre": "string",
            "uf_principal": "string",
            "setor_cnae_classe_principal": "string",
            "setor_cnae_subclasse_principal": "string",
            "subsetor_ibge_principal_codigo": "string",
        },
        chunksize=1_000_000,
        low_memory=False,
    )
    for chunk in chunks:
        total_rows += len(chunk)
        chunk = chunk[chunk["ano"].isin(use_years)].copy()
        if chunk.empty:
            continue
        chunk["cnpj_raiz"] = (
            chunk["cnpj_raiz"].astype(str).str.replace(r"\D", "", regex=True).str.zfill(8)
        )
        companies_seen.update(chunk["cnpj_raiz"].unique().tolist())
        chunk["variacao_pct_empregados"] = pd.to_numeric(
            chunk["variacao_pct_empregados"], errors="coerce"
        )
        chunk = chunk[chunk["variacao_pct_empregados"].notna()].copy()
        if chunk.empty:
            continue
        chunk["total_empregados"] = pd.to_numeric(chunk["total_empregados"], errors="coerce")
        chunk["empregados_lag"] = pd.to_numeric(chunk["empregados_lag"], errors="coerce")
        chunk["regiao"] = chunk["uf_principal"].map(region_from_uf)
        chunk["porte"] = chunk["faixa_emprego_pre"].fillna("Sem informacao")
        chunk["setor"] = (
            "Subsetor IBGE "
            + chunk["subsetor_ibge_principal_codigo"].fillna("").astype(str).str.zfill(2)
        )
        is_fust = chunk["cnpj_raiz"].isin(fust_roots)
        if is_fust.any():
            treated_parts.append(chunk[is_fust].copy())
        if (~is_fust).any():
            control_parts.append(chunk[~is_fust].copy())

    treated = pd.concat(treated_parts, ignore_index=True) if treated_parts else pd.DataFrame()
    controls = pd.concat(control_parts, ignore_index=True) if control_parts else pd.DataFrame()
    if not treated.empty:
        treated = treated.merge(
            profile_root[
                [
                    "cnpj_raiz",
                    "cnpj_key",
                    "ano_primeira_operacao",
                    "tipo_primeira_operacao",
                    "regiao_fust",
                    "porte_fust",
                    "setor_cnae_fust",
                    "subsetor_cnae_fust",
                ]
            ],
            on="cnpj_raiz",
            how="left",
        )
        treated = treated[treated["ano_primeira_operacao"].notna()].copy()
        treated["ano_primeira_operacao"] = treated["ano_primeira_operacao"].astype(int)

    info = {
        "arquivo_rais_universo": str(universo_path),
        "linhas_rais_universo": total_rows,
        "empresas_rais_universo_anos_usados": len(companies_seen),
        "empresas_fust_com_cnpj_raiz": len(fust_roots),
        "empresas_fust_sem_cnpj_raiz": missing_roots,
        "linhas_tratadas_com_variacao": len(treated),
        "linhas_controle_com_variacao": len(controls),
    }
    return treated, controls, info


def parse_painel_matching_rais() -> tuple[pd.DataFrame, dict[str, object]]:
    matched_path = DATA_DIR / "RAIS_matching_tratadas_controles_2018_2024.xlsx"
    if not matched_path.exists():
        matched_path = DATA_DIR / "RAIS_matching_tratadas_controles_2018_2024.csv"
    if not matched_path.exists():
        raise FileNotFoundError(
            f"{matched_path} nao encontrado. Rode gerar_painel_matching_rais.py antes."
        )
    if matched_path.suffix.lower() == ".xlsx":
        painel = pd.read_excel(matched_path, sheet_name=0, dtype="string")
    else:
        painel = pd.read_csv(matched_path, sep=";", dtype="string", low_memory=False)
    painel["ano"] = pd.to_numeric(painel["ano"], errors="coerce").astype("Int64")
    painel["cnpj_raiz"] = (
        painel["cnpj_raiz"].astype(str).str.replace(r"\D", "", regex=True).str.zfill(8)
    )

    monthly_path = DATA_DIR / "RAIS_matching_emprego_mensal_2022_2024.xlsx"
    if monthly_path.exists():
        monthly = pd.read_excel(monthly_path, sheet_name="emprego_mensal", dtype={"cnpj_raiz": "string"})
        monthly["cnpj_raiz"] = (
            monthly["cnpj_raiz"].astype(str).str.replace(r"\D", "", regex=True).str.zfill(8)
        )
        monthly["ano"] = pd.to_numeric(monthly["ano"], errors="coerce").astype("Int64")
        monthly["estoque_medio_anual_empregados"] = pd.to_numeric(
            monthly["estoque_medio_anual_empregados"], errors="coerce"
        )
        lookup = monthly[["cnpj_raiz", "ano", "estoque_medio_anual_empregados"]].drop_duplicates()
        painel = painel.merge(
            lookup.rename(columns={"estoque_medio_anual_empregados": "empregados_medio_ano_apoio"}),
            on=["cnpj_raiz", "ano"],
            how="left",
        )
        painel["_ano_pre"] = painel["ano"] - 1
        painel = painel.merge(
            lookup.rename(
                columns={
                    "ano": "_ano_pre",
                    "estoque_medio_anual_empregados": "empregados_medio_ano_anterior",
                }
            ),
            on=["cnpj_raiz", "_ano_pre"],
            how="left",
        )
        painel["_ano_pos"] = painel["ano"] + 1
        painel = painel.merge(
            lookup.rename(
                columns={
                    "ano": "_ano_pos",
                    "estoque_medio_anual_empregados": "empregados_medio_ano_posterior",
                }
            ),
            on=["cnpj_raiz", "_ano_pos"],
            how="left",
        )
        painel.drop(columns=["_ano_pre", "_ano_pos"], inplace=True)
        painel["variacao_pct_empregados_apoio"] = (
            painel["empregados_medio_ano_apoio"] - painel["empregados_medio_ano_anterior"]
        ) / painel["empregados_medio_ano_anterior"]
        painel.loc[
            painel["empregados_medio_ano_anterior"].isna()
            | (painel["empregados_medio_ano_anterior"] <= 0),
            "variacao_pct_empregados_apoio",
        ] = pd.NA
        painel["variacao_pct_empregados_pos"] = (
            painel["empregados_medio_ano_posterior"] - painel["empregados_medio_ano_apoio"]
        ) / painel["empregados_medio_ano_apoio"]
        painel.loc[
            painel["empregados_medio_ano_apoio"].isna()
            | (painel["empregados_medio_ano_apoio"] <= 0),
            "variacao_pct_empregados_pos",
        ] = pd.NA
        painel["faixa_emprego_pre_medio"] = painel[
            "empregados_medio_ano_anterior"
        ].map(employment_band)
        painel["distancia_emprego_medio_pre"] = pd.NA
        refined_parts = []
        for _, group in painel.groupby("request_id", dropna=False):
            treated = group[group["papel"].eq("tratada")].copy()
            controls = group[group["papel"].eq("controle")].copy()
            if treated.empty:
                refined_parts.append(group)
                continue
            treated_pre = treated["empregados_medio_ano_anterior"].iloc[0]
            treated_band = treated["faixa_emprego_pre_medio"].iloc[0]
            if pd.isna(treated_pre):
                refined_parts.append(pd.concat([treated, controls], ignore_index=True, sort=False))
                continue
            controls = controls[controls["empregados_medio_ano_anterior"].notna()].copy()
            same_band = controls[controls["faixa_emprego_pre_medio"].eq(treated_band)].copy()
            controls = same_band if not same_band.empty else controls
            controls["distancia_emprego_medio_pre"] = (
                controls["empregados_medio_ano_anterior"] - treated_pre
            ).abs()
            treated["distancia_emprego_medio_pre"] = 0
            controls = controls.sort_values(
                ["distancia_emprego_medio_pre", "cnpj_raiz"], kind="mergesort"
            ).drop_duplicates(["request_id", "cnpj_raiz"]).head(20)
            refined_parts.append(pd.concat([treated, controls], ignore_index=True, sort=False))
        painel = pd.concat(refined_parts, ignore_index=True, sort=False)
    else:
        for col in [
            "empregados_medio_ano_anterior",
            "empregados_medio_ano_apoio",
            "empregados_medio_ano_posterior",
            "variacao_pct_empregados_apoio",
            "variacao_pct_empregados_pos",
            "faixa_emprego_pre_medio",
            "distancia_emprego_medio_pre",
        ]:
            painel[col] = pd.NA

    for col in [
        "total_empregados",
        "empregados_lag",
        "variacao_pct_empregados",
        "empregados_medio_ano_anterior",
        "empregados_medio_ano_apoio",
        "empregados_medio_ano_posterior",
        "variacao_pct_empregados_apoio",
        "variacao_pct_empregados_pos",
        "distancia_emprego_pre",
        "distancia_emprego_medio_pre",
    ]:
        painel[col] = pd.to_numeric(painel[col], errors="coerce")
    info = {
        "arquivo_painel_matching": str(matched_path),
        "arquivo_rais_emprego_mensal": str(monthly_path) if monthly_path.exists() else "nao encontrado",
        "linhas_painel_matching": len(painel),
        "empresas_tratadas_matching": painel.loc[painel["papel"].eq("tratada"), "cnpj_raiz"].nunique(),
        "empresas_controle_matching": painel.loc[painel["papel"].eq("controle"), "cnpj_raiz"].nunique(),
        "requests_matching": painel["request_id"].nunique(),
        "anos_matching": ", ".join(map(str, sorted(painel["ano"].dropna().astype(int).unique()))),
    }
    return painel, info


def full_rais_panel(
    profile: pd.DataFrame,
    matched_panel: pd.DataFrame,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    monthly_path = DATA_DIR / "RAIS_matching_emprego_mensal_2022_2024.xlsx"
    if not monthly_path.exists():
        raise FileNotFoundError(
            f"{monthly_path} nao encontrado. Rode reprocess_rais_matching_emprego_mensal.R antes."
        )
    monthly = pd.read_excel(
        monthly_path,
        sheet_name="emprego_mensal_cnpj",
        dtype={"cnpj_raiz": "string", "cnpj_completo": "string"},
    )
    monthly["cnpj_raiz"] = (
        monthly["cnpj_raiz"].astype(str).str.replace(r"\D", "", regex=True).str.zfill(8)
    )
    monthly["cnpj_completo"] = (
        monthly["cnpj_completo"].astype(str).str.replace(r"\D", "", regex=True).str.zfill(14)
    )
    monthly["ano"] = pd.to_numeric(monthly["ano"], errors="coerce").astype("Int64")
    monthly["estoque_medio_anual_empregados"] = pd.to_numeric(
        monthly["estoque_medio_anual_empregados"], errors="coerce"
    )

    key_root = cnpj_key_to_root_map()
    key_full = cnpj_key_to_full_map()
    profile_root = profile.merge(key_root, on="cnpj_key", how="left").merge(
        key_full, on="cnpj_key", how="left"
    )
    profile_root["cnpj_raiz"] = profile_root["cnpj_raiz"].fillna(profile_root["cnpj_raiz_base"])
    profile_root["cnpj_completo"] = profile_root["cnpj_completo_base"].fillna(
        profile_root["cnpj_completo"]
    )
    profile_root = profile_root[profile_root["cnpj_raiz"].notna()].copy()
    profile_root["cnpj_raiz"] = (
        profile_root["cnpj_raiz"].astype(str).str.replace(r"\D", "", regex=True).str.zfill(8)
    )
    profile_root["cnpj_completo"] = (
        profile_root["cnpj_completo"].astype(str).str.replace(r"\D", "", regex=True).str.zfill(14)
    )
    profile_root.loc[profile_root["cnpj_completo"].eq("00000000000000"), "cnpj_completo"] = pd.NA
    profile_root = profile_root[profile_root["cnpj_completo"].notna()].copy()
    treated_meta = (
        profile_root.sort_values(["cnpj_completo", "ano_primeira_operacao", "data_primeira_operacao"])
        .groupby("cnpj_completo", dropna=False)
        .agg(
            cnpj_raiz=("cnpj_raiz", first_non_empty),
            cnpj_key=("cnpj_key", first_non_empty),
            cliente=("cliente", first_non_empty),
            ano_primeira_operacao=("ano_primeira_operacao", "min"),
            tipo_primeira_operacao=("tipo_primeira_operacao", first_non_empty),
            regiao=("regiao_fust", first_non_empty),
            porte_fust=("porte_fust", first_non_empty),
            setor_fust=("subsetor_cnae_fust", first_non_empty),
            valor_operacao_fust=("valor_operacao_fust", "sum"),
            qtd_contratos_fust=("qtd_contratos_fust", "sum"),
        )
        .reset_index()
    )
    treated = monthly.merge(
        treated_meta,
        on=["cnpj_raiz", "cnpj_completo"],
        how="inner",
    )
    treated["papel"] = "tratada"
    treated["tipo_operacao"] = treated["tipo_primeira_operacao"].fillna("Nao informado")
    treated["setor"] = treated["setor_fust"].fillna("Nao informado")

    controls_meta = (
        matched_panel[matched_panel["papel"].eq("controle")]
        .sort_values(["cnpj_raiz", "ano"])
        .groupby("cnpj_raiz", dropna=False)
        .agg(
            regiao=("regiao", first_non_empty),
            setor=("setor", first_non_empty),
        )
        .reset_index()
    )
    controls = monthly.merge(controls_meta, on="cnpj_raiz", how="inner")
    main_controls = (
        controls.groupby(["cnpj_raiz", "cnpj_completo"], dropna=False)
        .agg(estoque_medio_referencia=("estoque_medio_anual_empregados", "mean"))
        .reset_index()
        .sort_values(["cnpj_raiz", "estoque_medio_referencia", "cnpj_completo"], ascending=[True, False, True])
        .drop_duplicates("cnpj_raiz")
    )
    controls = controls.merge(
        main_controls[["cnpj_raiz", "cnpj_completo"]],
        on=["cnpj_raiz", "cnpj_completo"],
        how="inner",
    )
    controls["papel"] = "controle"
    controls["cnpj_key"] = pd.NA
    controls["cliente"] = pd.NA
    controls["ano_primeira_operacao"] = pd.NA
    controls["tipo_primeira_operacao"] = pd.NA
    controls["tipo_operacao"] = "Controle"
    controls["porte_fust"] = pd.NA
    controls["valor_operacao_fust"] = pd.NA
    controls["qtd_contratos_fust"] = pd.NA

    panel = pd.concat([treated, controls], ignore_index=True, sort=False)
    panel = panel.sort_values(["papel", "cnpj_completo", "ano"])
    panel["estoque_medio_lag"] = panel.groupby(["papel", "cnpj_completo"])[
        "estoque_medio_anual_empregados"
    ].shift(1)
    panel["ano_lag"] = panel.groupby(["papel", "cnpj_completo"])["ano"].shift(1)
    panel["variacao_pct_empregados"] = (
        panel["estoque_medio_anual_empregados"] - panel["estoque_medio_lag"]
    ) / panel["estoque_medio_lag"]
    panel.loc[
        panel["estoque_medio_lag"].isna()
        | (panel["estoque_medio_lag"] <= 0)
        | (panel["ano_lag"] != panel["ano"] - 1),
        "variacao_pct_empregados",
    ] = pd.NA
    panel["estoque_medio_lead"] = panel.groupby(["papel", "cnpj_completo"])[
        "estoque_medio_anual_empregados"
    ].shift(-1)
    panel["ano_lead"] = panel.groupby(["papel", "cnpj_completo"])["ano"].shift(-1)
    panel["variacao_pct_empregados_posterior"] = (
        panel["estoque_medio_lead"] - panel["estoque_medio_anual_empregados"]
    ) / panel["estoque_medio_anual_empregados"]
    panel.loc[
        panel["estoque_medio_anual_empregados"].isna()
        | (panel["estoque_medio_anual_empregados"] <= 0)
        | (panel["ano_lead"] != panel["ano"] + 1),
        "variacao_pct_empregados_posterior",
    ] = pd.NA
    panel["porte"] = panel["estoque_medio_lag"].map(employment_band)
    panel["regiao"] = panel["regiao"].fillna("Nao informado")
    panel["setor"] = panel["setor"].fillna("Nao informado")

    painel_empresas = panel[
        [
            "papel",
            "cnpj_raiz",
            "cnpj_completo",
            "cnpj_key",
            "cliente",
            "ano",
            "ano_primeira_operacao",
            "tipo_operacao",
            "regiao",
            "porte",
            "setor",
            "estoque_medio_lag",
            "estoque_medio_anual_empregados",
            "estoque_medio_lead",
            "variacao_pct_empregados",
            "variacao_pct_empregados_posterior",
            "valor_operacao_fust",
            "qtd_contratos_fust",
        ]
    ].sort_values(["papel", "cnpj_completo", "ano"])
    return panel, painel_empresas


def efficacy_rows(base: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    contracts = contract_view(base)
    years = sorted(contracts["ano_aprovacao"].dropna().astype(int).unique())
    tipos = sorted(contracts["tipo_operacao"].dropna().unique())
    tipos_out = [*tipos, "Total"]
    first_year = contracts.groupby("cnpj_key")["ano_aprovacao"].min()
    dimensions = [
        ("Geral", None, ["Total"]),
        ("Regiao", "regiao", None),
        ("Porte", "porte", None),
        ("Setor", "setor", None),
    ]

    for year in years:
        for tipo in tipos_out:
            if tipo == "Total":
                base_subset = contracts[contracts["ano_aprovacao"] == year]
            else:
                base_subset = contracts[
                    (contracts["ano_aprovacao"] == year) & (contracts["tipo_operacao"] == tipo)
                ]
            for dimension, category_col, fixed_categories in dimensions:
                if fixed_categories is not None:
                    categories = fixed_categories
                else:
                    categories = sorted(
                        base_subset[category_col]
                        .fillna("Nao informado")
                        .astype(str)
                        .unique()
                        .tolist()
                    )

                for category in categories:
                    if category_col is None:
                        subset = base_subset
                    else:
                        subset = base_subset[
                            base_subset[category_col].fillna("Nao informado").astype(str)
                            == category
                        ]
                    new_clients = subset[
                        subset["cnpj_key"].map(first_year).astype("Int64") == year
                    ]["cnpj_key"].nunique()
                    metrics = [
                        (
                            "I1",
                            "numero de operacoes aprovadas",
                            subset["contrato_id"].nunique(),
                            subset["contrato_id"].nunique(),
                            subset["cnpj_key"].nunique(),
                            "Contratos contados por ano de aprovacao; quando o identificador esta ausente, cada linha da base operacional e tratada como contrato.",
                        ),
                        (
                            "I2",
                            "numero de clientes apoiados",
                            subset["cnpj_key"].nunique(),
                            subset["contrato_id"].nunique(),
                            subset["cnpj_key"].nunique(),
                            "Clientes contados por CPF/CNPJ normalizado e mascarado.",
                        ),
                        (
                            "I3",
                            "novos clientes",
                            new_clients,
                            subset["contrato_id"].nunique(),
                            subset["cnpj_key"].nunique(),
                            "Proxy de primeira aparicao na base FUST compilada; a base nao contem todo o historico BNDES.",
                        ),
                    ]
                    for codigo, indicador, valor, operacoes, clientes, observacao in metrics:
                        rows.append(
                            {
                                "nivel_resultado": "Eficacia",
                                "codigo_indicador": codigo,
                                "indicador": indicador,
                                "estatistica": "valor observado",
                                "tipo_operacao": tipo,
                                "ano_aprovacao": year,
                                "janela_calculo": "nao aplicavel",
                                "dimensao": dimension,
                                "categoria": category,
                                "valor_indicador": valor,
                                "registros": subset["linhas_base"].sum(),
                                "operacoes": operacoes,
                                "clientes": clientes,
                                "controles_distintos": pd.NA,
                                "matches_tratado_controle": pd.NA,
                                "observacao": observacao,
                            }
                        )

    for year in years:
        for tipo in tipos_out:
            for codigo, indicador in [
                ("I4", "equipamentos adquiridos ou modernizados por tipo"),
                ("I5", "area construida ou modernizada em projetos de investimento"),
            ]:
                rows.append(
                    {
                        "nivel_resultado": "Eficacia",
                        "codigo_indicador": codigo,
                        "indicador": indicador,
                        "estatistica": "nao disponivel",
                        "tipo_operacao": tipo,
                        "ano_aprovacao": year,
                        "janela_calculo": "nao aplicavel",
                        "dimensao": "Geral",
                        "categoria": "Total",
                        "valor_indicador": pd.NA,
                        "registros": pd.NA,
                        "operacoes": pd.NA,
                        "clientes": pd.NA,
                        "controles_distintos": pd.NA,
                        "matches_tratado_controle": pd.NA,
                        "observacao": "Indicador nao observado diretamente na Base de Operacoes_FUST_compilada nem no Painel RAIS.",
                    }
                )

    return pd.DataFrame(rows)


def matched_control_values(
    treated: pd.DataFrame,
    controls: pd.DataFrame,
    category_col: str | None,
) -> tuple[list[float], set[str], int]:
    values: list[float] = []
    control_keys: set[str] = set()
    match_count = 0
    for _, row in treated.iterrows():
        pool = controls[controls["faixa_emprego_pre"] == row["faixa_emprego_pre"]]
        if category_col is not None:
            pool = pool[pool[category_col] == row[category_col]]
        pool = pool[pool["variacao_pct_empregados"].notna()]
        if pool.empty:
            continue
        vals = pool["variacao_pct_empregados"].astype(float).tolist()
        values.extend(vals)
        key_col = "cnpj_key" if "cnpj_key" in pool.columns else "cnpj_raiz"
        control_keys.update(pool[key_col].tolist())
        match_count += len(pool)
    return values, control_keys, match_count


def effectiveness_rows(
    treated_universe: pd.DataFrame,
    control_universe: pd.DataFrame,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    rows: list[dict[str, object]] = []
    match_audit: list[dict[str, object]] = []
    years = sorted(treated_universe["ano"].dropna().astype(int).unique())
    tipos = sorted(treated_universe["tipo_primeira_operacao"].dropna().unique())
    dimensions = [
        ("Geral", None, ["Total"]),
        ("Regiao", "regiao", None),
        ("Porte", "porte", None),
        ("Setor", "setor", None),
    ]

    for year in years:
        year_treated = treated_universe[
            (treated_universe["ano"] == year)
            & (treated_universe["variacao_pct_empregados"].notna())
            & (treated_universe["ano_primeira_operacao"].notna())
        ].copy()
        controls_base = control_universe[
            (control_universe["ano"] == year)
            & (control_universe["variacao_pct_empregados"].notna())
        ].copy()
        for tipo in tipos:
            treated_base = year_treated[
                (year_treated["ano_primeira_operacao"] == year)
                & (year_treated["tipo_primeira_operacao"] == tipo)
            ].copy()

            for dimension, category_col, fixed_categories in dimensions:
                if fixed_categories is not None:
                    categories = fixed_categories
                else:
                    categories = sorted(
                        treated_base[category_col]
                        .fillna("Nao informado")
                        .astype(str)
                        .unique()
                        .tolist()
                    )

                for category in categories:
                    if category_col is None:
                        treated = treated_base
                    else:
                        treated = treated_base[
                            treated_base[category_col].fillna("Nao informado").astype(str)
                            == category
                        ]
                    treated_values = treated["variacao_pct_empregados"].astype(float).tolist()
                    control_values, control_keys, match_count = matched_control_values(
                        treated, controls_base, category_col
                    )

                    for stat_name, stat_func in [
                        ("mediana", safe_median),
                        ("media", safe_mean),
                    ]:
                        i6_value = stat_func(treated_values)
                        i7_value = stat_func(control_values)
                        rows.append(
                            {
                                "nivel_resultado": "Efetividade",
                                "codigo_indicador": "I6",
                                "indicador": "variacao percentual do estoque medio anual de empregados das empresas apoiadas",
                                "estatistica": stat_name,
                                "tipo_operacao": tipo,
                                "ano_aprovacao": year,
                                "dimensao": dimension,
                                "categoria": category,
                                "valor_indicador": i6_value,
                                "registros": len(treated),
                                "operacoes": pd.NA,
                                "clientes": treated["cnpj_key"].nunique(),
                                "controles_distintos": len(control_keys),
                                "matches_tratado_controle": match_count,
                                "observacao": "Coorte de empresas no primeiro ano de apoio FUST com RAIS em ano-1 e ano.",
                            }
                        )
                        rows.append(
                            {
                                "nivel_resultado": "Efetividade",
                                "codigo_indicador": "I7",
                                "indicador": "indicador de controle para comparacao",
                                "estatistica": stat_name,
                                "tipo_operacao": tipo,
                                "ano_aprovacao": year,
                                "dimensao": dimension,
                                "categoria": category,
                                "valor_indicador": i7_value,
                                "registros": len(treated),
                                "operacoes": pd.NA,
                                "clientes": treated["cnpj_key"].nunique(),
                                "controles_distintos": len(control_keys),
                                "matches_tratado_controle": match_count,
                                "observacao": "Controle composto por empresas da RAIS reprocessada que nao constam na base FUST, pareadas por faixa de emprego pre-apoio"
                                + (" e pela dimensao da tabela." if category_col else "."),
                            }
                        )

                    match_audit.append(
                        {
                            "ano_aprovacao": year,
                            "tipo_operacao": tipo,
                            "dimensao": dimension,
                            "categoria": category,
                            "tratados_com_variacao": treated["cnpj_key"].nunique(),
                            "controles_distintos": len(control_keys),
                            "matches_tratado_controle": match_count,
                            "faixas_tratadas": ", ".join(
                                sorted(treated["faixa_emprego_pre"].dropna().unique())
                            ),
                        }
                    )

    return pd.DataFrame(rows), pd.DataFrame(match_audit)


def effectiveness_rows_matched(painel: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    rows: list[dict[str, object]] = []
    match_audit: list[dict[str, object]] = []
    treated_panel = painel[painel["papel"].eq("tratada")].copy()
    controls_panel = painel[painel["papel"].eq("controle")].copy()
    windows = [
        (
            "ano_apoio_vs_anterior",
            "ano do apoio vs ano anterior",
            "variacao_pct_empregados_apoio",
        ),
        (
            "ano_posterior_vs_apoio",
            "ano posterior vs ano do apoio",
            "variacao_pct_empregados_pos",
        ),
    ]

    grouped_items = []
    group_cols = ["ano", "tipo_primeira_operacao", "dimensao", "categoria"]
    for keys, treated in treated_panel.groupby(group_cols, dropna=False):
        year, tipo, dimension, category = keys
        grouped_items.append((year, tipo, dimension, category, treated.copy()))
    total_group_cols = ["ano", "dimensao", "categoria"]
    for keys, treated in treated_panel.groupby(total_group_cols, dropna=False):
        year, dimension, category = keys
        grouped_items.append((year, "Total", dimension, category, treated.copy()))

    for year, tipo, dimension, category, treated in grouped_items:
        request_ids = set(treated["request_id"])
        controls = controls_panel[controls_panel["request_id"].isin(request_ids)].copy()

        for window_code, window_label, value_col in windows:
            treated_valid = treated[treated[value_col].notna()].copy()
            controls_valid = (
                controls[controls[value_col].notna()]
                .sort_values(["cnpj_raiz", "distancia_emprego_medio_pre"], kind="mergesort")
                .drop_duplicates("cnpj_raiz")
                .copy()
            )
            treated_values = treated_valid[value_col].astype(float).tolist()
            control_values = controls_valid[value_col].astype(float).tolist()

            for stat_name, stat_func in [("mediana", safe_median), ("media", safe_mean)]:
                rows.append(
                    {
                        "nivel_resultado": "Efetividade",
                        "codigo_indicador": "I6",
                        "indicador": "variacao percentual do estoque medio anual de empregados das empresas apoiadas",
                        "estatistica": stat_name,
                        "tipo_operacao": tipo,
                        "ano_aprovacao": int(year),
                        "janela_calculo": window_label,
                        "dimensao": dimension,
                        "categoria": category,
                        "valor_indicador": stat_func(treated_values),
                        "registros": len(treated_valid),
                        "operacoes": pd.NA,
                        "clientes": treated_valid["cnpj_raiz"].nunique(),
                        "controles_distintos": controls_valid["cnpj_raiz"].nunique(),
                        "matches_tratado_controle": len(controls_valid),
                        "observacao": "Variacao calculada com estoque medio mensal anual da RAIS para a janela: "
                        + window_label
                        + ".",
                    }
                )
                rows.append(
                    {
                        "nivel_resultado": "Efetividade",
                        "codigo_indicador": "I7",
                        "indicador": "indicador de controle para comparacao",
                        "estatistica": stat_name,
                        "tipo_operacao": tipo,
                        "ano_aprovacao": int(year),
                        "janela_calculo": window_label,
                        "dimensao": dimension,
                        "categoria": category,
                        "valor_indicador": stat_func(control_values),
                        "registros": len(treated_valid),
                        "operacoes": pd.NA,
                        "clientes": treated_valid["cnpj_raiz"].nunique(),
                        "controles_distintos": controls_valid["cnpj_raiz"].nunique(),
                        "matches_tratado_controle": len(controls_valid),
                        "observacao": "Controle composto pelo painel RAIS pos-matching: ate 20 controles por tratada, pareados por ano, faixa de emprego pre-apoio e dimensao da tabela; variacao calculada com estoque medio mensal anual.",
                    }
                )

            match_audit.append(
                {
                    "ano_aprovacao": int(year),
                    "tipo_operacao": tipo,
                    "janela_calculo": window_label,
                    "dimensao": dimension,
                    "categoria": category,
                    "tratados_com_variacao": treated_valid["cnpj_raiz"].nunique(),
                    "controles_distintos": controls_valid["cnpj_raiz"].nunique(),
                    "matches_tratado_controle": len(controls_valid),
                    "tratadas_media": safe_mean(treated_values),
                    "tratadas_mediana": safe_median(treated_values),
                    "controles_media": safe_mean(control_values),
                    "controles_mediana": safe_median(control_values),
                    "controles_p95": safe_quantile(control_values, 0.95),
                    "controles_max": max(control_values) if control_values else math.nan,
                    "faixas_tratadas": ", ".join(
                        sorted(treated["faixa_emprego_pre"].dropna().astype(str).unique())
                    ),
                }
            )

    return pd.DataFrame(rows), pd.DataFrame(match_audit)


def effectiveness_rows_full_panel(panel: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    rows: list[dict[str, object]] = []
    audit: list[dict[str, object]] = []
    treated_panel = panel[panel["papel"].eq("tratada")].copy()
    controls_panel = panel[panel["papel"].eq("controle")].copy()
    windows = [
        (
            "ano de referencia vs anterior",
            "variacao_pct_empregados",
            "crescimento do estoque medio anual em relacao ao ano anterior",
        ),
        (
            "ano posterior vs referencia",
            "variacao_pct_empregados_posterior",
            "crescimento do estoque medio anual do ano posterior em relacao ao ano de referencia",
        ),
    ]

    years = sorted(treated_panel["ano"].dropna().astype(int).unique())
    tipos = sorted(treated_panel["tipo_operacao"].dropna().unique())
    dimensions = [
        ("Geral", None, ["Total"]),
        ("Regiao", "regiao", None),
        ("Porte", "porte", None),
        ("Setor", "setor", None),
    ]

    for year in years:
        treated_year = treated_panel[treated_panel["ano"].astype(int).eq(year)].copy()
        controls_year = controls_panel[controls_panel["ano"].astype(int).eq(year)].copy()
        for tipo in [*tipos, "Total"]:
            treated_base = (
                treated_year
                if tipo == "Total"
                else treated_year[treated_year["tipo_operacao"].eq(tipo)].copy()
            )
            for dimension, category_col, fixed_categories in dimensions:
                if fixed_categories is not None:
                    categories = fixed_categories
                else:
                    categories = sorted(
                        treated_base[category_col]
                        .fillna("Nao informado")
                        .astype(str)
                        .unique()
                        .tolist()
                    )

                for category in categories:
                    if category_col is None:
                        treated = treated_base
                        controls = controls_year
                    else:
                        treated = treated_base[
                            treated_base[category_col].fillna("Nao informado").astype(str)
                            == category
                        ]
                        controls = controls_year[
                            controls_year[category_col].fillna("Nao informado").astype(str)
                            == category
                        ]

                    for window_label, value_col, window_note in windows:
                        treated_valid = treated[treated[value_col].notna()].copy()
                        controls_valid = controls[controls[value_col].notna()].copy()
                        treated_values = treated_valid[value_col].astype(float).tolist()
                        control_values = controls_valid[value_col].astype(float).tolist()
                        for stat_name, stat_func in [("mediana", safe_median), ("media", safe_mean)]:
                            rows.append(
                                {
                                    "nivel_resultado": "Efetividade",
                                    "codigo_indicador": "I6",
                                    "indicador": "variacao percentual do estoque medio anual de empregados das empresas apoiadas",
                                    "estatistica": stat_name,
                                    "tipo_operacao": tipo,
                                    "ano_aprovacao": int(year),
                                    "janela_calculo": window_label,
                                    "dimensao": dimension,
                                    "categoria": category,
                                    "valor_indicador": stat_func(treated_values),
                                    "registros": len(treated_valid),
                                    "operacoes": pd.NA,
                                    "clientes": treated_valid["cnpj_completo"].nunique(),
                                    "controles_distintos": controls_valid["cnpj_completo"].nunique(),
                                    "matches_tratado_controle": len(controls_valid),
                                    "observacao": "Painel completo de empresas apoiadas FUST mapeadas na RAIS, independente do ano de apoio; "
                                    + window_note
                                    + ".",
                                }
                            )
                            rows.append(
                                {
                                    "nivel_resultado": "Efetividade",
                                    "codigo_indicador": "I7",
                                    "indicador": "indicador de controle para comparacao",
                                    "estatistica": stat_name,
                                    "tipo_operacao": tipo,
                                    "ano_aprovacao": int(year),
                                    "janela_calculo": window_label,
                                    "dimensao": dimension,
                                    "categoria": category,
                                    "valor_indicador": stat_func(control_values),
                                    "registros": len(treated_valid),
                                    "operacoes": pd.NA,
                                    "clientes": treated_valid["cnpj_completo"].nunique(),
                                    "controles_distintos": controls_valid["cnpj_completo"].nunique(),
                                    "matches_tratado_controle": len(controls_valid),
                                    "observacao": "Controle formado pelas empresas nao apoiadas selecionadas no matching e acompanhadas no painel RAIS anual; "
                                    + window_note
                                    + ".",
                                }
                            )

                        audit.append(
                            {
                                "ano_aprovacao": int(year),
                                "tipo_operacao": tipo,
                                "janela_calculo": window_label,
                                "dimensao": dimension,
                                "categoria": category,
                                "tratados_com_variacao": treated_valid["cnpj_completo"].nunique(),
                                "controles_distintos": controls_valid["cnpj_completo"].nunique(),
                                "matches_tratado_controle": len(controls_valid),
                                "tratadas_media": safe_mean(treated_values),
                                "tratadas_mediana": safe_median(treated_values),
                                "controles_media": safe_mean(control_values),
                                "controles_mediana": safe_median(control_values),
                                "controles_p95": safe_quantile(control_values, 0.95),
                                "controles_max": max(control_values) if control_values else math.nan,
                                "faixas_tratadas": ", ".join(
                                    sorted(treated_valid["porte"].dropna().astype(str).unique())
                                ),
                            }
                        )

    return pd.DataFrame(rows), pd.DataFrame(audit)


def summarize_by_year(all_rows: pd.DataFrame) -> pd.DataFrame:
    return (
        all_rows.groupby(["ano_aprovacao", "nivel_resultado"], dropna=False)
        .agg(
            indicadores_distintos=("codigo_indicador", "nunique"),
            linhas=("codigo_indicador", "size"),
            tipos_operacao=("tipo_operacao", "nunique"),
            dimensoes=("dimensao", "nunique"),
            observacao=("observacao", first_non_empty),
        )
        .reset_index()
        .sort_values(["ano_aprovacao", "nivel_resultado"])
    )


def coverage_rais_by_year_type(all_rows: pd.DataFrame) -> pd.DataFrame:
    efficacy_clients = all_rows[
        (all_rows["codigo_indicador"] == "I2")
        & (all_rows["dimensao"] == "Geral")
        & (all_rows["categoria"] == "Total")
    ][["ano_aprovacao", "tipo_operacao", "valor_indicador"]].rename(
        columns={"valor_indicador": "clientes_eficacia_i2"}
    )
    treated_rais = all_rows[
        (all_rows["codigo_indicador"] == "I6")
        & (all_rows["estatistica"] == "mediana")
        & (all_rows["janela_calculo"] == "ano de referencia vs anterior")
        & (all_rows["dimensao"] == "Geral")
        & (all_rows["categoria"] == "Total")
    ][["ano_aprovacao", "tipo_operacao", "clientes", "controles_distintos"]].rename(
        columns={
            "clientes": "empresas_tratadas_com_rais_i6",
            "controles_distintos": "controles_distintos_i7",
        }
    )
    coverage = efficacy_clients.merge(
        treated_rais, on=["ano_aprovacao", "tipo_operacao"], how="left"
    )
    coverage["empresas_tratadas_com_rais_i6"] = coverage[
        "empresas_tratadas_com_rais_i6"
    ].fillna(0)
    coverage["controles_distintos_i7"] = coverage["controles_distintos_i7"].fillna(0)
    for col in [
        "clientes_eficacia_i2",
        "empresas_tratadas_com_rais_i6",
        "controles_distintos_i7",
    ]:
        coverage[col] = pd.to_numeric(coverage[col], errors="coerce")
    coverage["observacao"] = (
        "I6/I7 usam o painel completo das empresas apoiadas, independente do ano de apoio. "
        "Por isso, empresas_tratadas_com_rais_i6 nao deve ser dividido pelos clientes apoiados no mesmo ano."
    )
    return coverage.sort_values(["ano_aprovacao", "tipo_operacao"])


def approval_date_coverage(base: pd.DataFrame) -> pd.DataFrame:
    contracts = contract_view(base)
    return (
        contracts.groupby(["tipo_operacao", "ano_aprovacao", "fonte_data_aprovacao"], dropna=False)
        .agg(
            operacoes=("contrato_id", "size"),
            clientes=("cnpj_key", "nunique"),
            primeira_data_aprovacao=("data_aprovacao", "min"),
            ultima_data_aprovacao=("data_aprovacao", "max"),
        )
        .reset_index()
        .sort_values(["tipo_operacao", "ano_aprovacao", "fonte_data_aprovacao"])
    )


def total_efficacy_rows(base: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    contracts = contract_view(base)
    first_year = contracts.groupby("cnpj_key")["ano_aprovacao"].min()
    period_years = sorted(contracts["ano_aprovacao"].dropna().astype(int).unique())
    tipos_out = [*sorted(contracts["tipo_operacao"].dropna().unique()), "Total"]
    for tipo in tipos_out:
        subset = contracts if tipo == "Total" else contracts[contracts["tipo_operacao"] == tipo]
        new_clients = 0
        for year in period_years:
            year_subset = subset[subset["ano_aprovacao"] == year]
            new_clients += year_subset[
                year_subset["cnpj_key"].map(first_year).astype("Int64") == year
            ]["cnpj_key"].nunique()
        metrics = [
            (
                "I1",
                "numero de operacoes aprovadas",
                len(subset),
                "Contratos-ano distintos no periodo.",
            ),
            (
                "I2",
                "numero de clientes apoiados",
                subset["cnpj_key"].nunique(),
                "Clientes distintos no periodo.",
            ),
            (
                "I3",
                "novos clientes",
                new_clients,
                "Clientes cuja primeira aparicao na base FUST compilada ocorre no periodo.",
            ),
        ]
        for codigo, indicador, valor, observacao in metrics:
            operacoes = valor if codigo == "I1" else len(subset)
            rows.append(
                {
                    "nivel_resultado": "Eficacia",
                    "codigo_indicador": codigo,
                    "indicador": indicador,
                    "estatistica": "valor observado",
                    "tipo_operacao": tipo,
                    "ano_aprovacao": "Total",
                    "janela_calculo": "nao aplicavel",
                    "dimensao": "Geral",
                    "categoria": "Total",
                    "valor_indicador": valor,
                    "registros": subset["linhas_base"].sum(),
                    "operacoes": operacoes,
                    "clientes": subset["cnpj_key"].nunique(),
                    "controles_distintos": pd.NA,
                    "matches_tratado_controle": pd.NA,
                    "observacao": observacao,
                }
            )

    return pd.DataFrame(rows)


def total_rows(base: pd.DataFrame, effectiveness: pd.DataFrame) -> pd.DataFrame:
    total_ef = total_efficacy_rows(base)

    eff = effectiveness[effectiveness["dimensao"] == "Geral"].copy()
    total_eff_rows = []
    for keys, group in eff.groupby(
        [
            "nivel_resultado",
            "codigo_indicador",
            "indicador",
            "estatistica",
            "tipo_operacao",
            "janela_calculo",
            "dimensao",
            "categoria",
        ],
        dropna=False,
    ):
        rec = dict(
            zip(
                [
                    "nivel_resultado",
                    "codigo_indicador",
                    "indicador",
                    "estatistica",
                    "tipo_operacao",
                    "janela_calculo",
                    "dimensao",
                    "categoria",
                ],
                keys,
            )
        )
        numeric_values = pd.to_numeric(group["valor_indicador"], errors="coerce")
        rec["valor_indicador"] = safe_mean(numeric_values.tolist())
        rec["registros"] = group["registros"].sum(skipna=True)
        rec["operacoes"] = pd.NA
        rec["clientes"] = group["clientes"].sum(skipna=True)
        rec["controles_distintos"] = group["controles_distintos"].sum(skipna=True)
        rec["matches_tratado_controle"] = group["matches_tratado_controle"].sum(skipna=True)
        rec["observacao"] = "Media dos valores anuais disponiveis na dimensao Geral."
        rec["ano_aprovacao"] = "Total"
        total_eff_rows.append(rec)
    total_eff = pd.DataFrame(total_eff_rows)
    return pd.concat([total_ef, total_eff], ignore_index=True, sort=False)[
        [
            "nivel_resultado",
            "codigo_indicador",
            "indicador",
            "estatistica",
            "tipo_operacao",
            "ano_aprovacao",
            "janela_calculo",
            "dimensao",
            "categoria",
            "valor_indicador",
            "registros",
            "operacoes",
            "clientes",
            "controles_distintos",
            "matches_tratado_controle",
            "observacao",
        ]
    ]


def add_missing_effectiveness_years(
    rows: pd.DataFrame, years: list[int], tipos: list[str]
) -> pd.DataFrame:
    available = set(
        tuple(x)
        for x in rows[["codigo_indicador", "tipo_operacao", "ano_aprovacao", "janela_calculo"]]
        .drop_duplicates()
        .itertuples(index=False, name=None)
    )
    windows = [
        "ano de referencia vs anterior",
        "ano posterior vs referencia",
    ]
    additions = []
    for year in years:
        for tipo in tipos:
            for codigo, indicador in [
                (
                    "I6",
                    "variacao percentual do estoque medio anual de empregados das empresas apoiadas",
                ),
                ("I7", "indicador de controle para comparacao"),
            ]:
                for window in windows:
                    if (codigo, tipo, year, window) in available:
                        continue
                    for stat in ["mediana", "media"]:
                        additions.append(
                            {
                                "nivel_resultado": "Efetividade",
                                "codigo_indicador": codigo,
                                "indicador": indicador,
                                "estatistica": stat,
                                "tipo_operacao": tipo,
                                "ano_aprovacao": year,
                                "janela_calculo": window,
                                "dimensao": "Geral",
                                "categoria": "Total",
                                "valor_indicador": pd.NA,
                                "registros": 0,
                                "operacoes": pd.NA,
                                "clientes": 0,
                                "controles_distintos": 0,
                                "matches_tratado_controle": 0,
                                "observacao": "Sem RAIS disponivel ou sem empresas apoiadas com variacao observavel neste ano.",
                            }
                        )
    if additions:
        rows = pd.concat([rows, pd.DataFrame(additions)], ignore_index=True, sort=False)
    return rows


def unavailable_effectiveness_rows(years: list[int], tipos: list[str]) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    for year in years:
        for tipo in tipos:
            for codigo, indicador in [
                ("I8", "variacao percentual do faturamento"),
                ("I9", "variacao percentual do investimento fixo das empresas apoiadas"),
            ]:
                rows.append(
                    {
                        "nivel_resultado": "Efetividade",
                        "codigo_indicador": codigo,
                        "indicador": indicador,
                        "estatistica": "nao disponivel",
                        "tipo_operacao": tipo,
                        "ano_aprovacao": year,
                        "janela_calculo": "nao aplicavel",
                        "dimensao": "Geral",
                        "categoria": "Total",
                        "valor_indicador": pd.NA,
                        "registros": pd.NA,
                        "operacoes": pd.NA,
                        "clientes": pd.NA,
                        "controles_distintos": pd.NA,
                        "matches_tratado_controle": pd.NA,
                        "observacao": "Indicador nao observado diretamente na Base de Operacoes_FUST_compilada nem no Painel RAIS.",
                    }
                )
    return pd.DataFrame(rows)


def save_outputs(
    all_rows: pd.DataFrame,
    summary: pd.DataFrame,
    match_audit: pd.DataFrame,
    painel_empresas_rais: pd.DataFrame,
    base: pd.DataFrame,
    rais_info: dict[str, object],
) -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    min_year = int(base["ano_aprovacao"].dropna().min())
    max_year = int(base["ano_aprovacao"].dropna().max())
    base_name = f"{TITULO_BASE} - {min_year} a {max_year}"
    xlsx_path = OUTPUT_DIR / f"{base_name}.xlsx"
    with open(xlsx_path, "a+b"):
        pass

    by_year = {
        str(year): all_rows[all_rows["ano_aprovacao"] == year].sort_values(
            [
                "nivel_resultado",
                "codigo_indicador",
                "tipo_operacao",
                "janela_calculo",
                "dimensao",
                "categoria",
                "estatistica",
            ]
        )
        for year in sorted(base["ano_aprovacao"].dropna().astype(int).unique())
    }
    sheets = {
        **by_year,
        "total": total_rows(base, all_rows[all_rows["nivel_resultado"] == "Efetividade"]),
        "por_ano": all_rows.sort_values(
            [
                "ano_aprovacao",
                "nivel_resultado",
                "codigo_indicador",
                "tipo_operacao",
                "janela_calculo",
                "dimensao",
                "categoria",
                "estatistica",
            ]
        ),
        "resumo_ano": summary,
        "eficacia": all_rows[all_rows["nivel_resultado"] == "Eficacia"],
        "efetividade": all_rows[all_rows["nivel_resultado"] == "Efetividade"],
        "efetividade_regiao": all_rows[
            (all_rows["nivel_resultado"] == "Efetividade") & (all_rows["dimensao"] == "Regiao")
        ],
        "efetividade_porte": all_rows[
            (all_rows["nivel_resultado"] == "Efetividade") & (all_rows["dimensao"] == "Porte")
        ],
        "efetividade_setor": all_rows[
            (all_rows["nivel_resultado"] == "Efetividade") & (all_rows["dimensao"] == "Setor")
        ],
        "auditoria_matching": match_audit,
        "cobertura_data_aprovacao": approval_date_coverage(base),
        "cobertura_rais_ano_tipo": coverage_rais_by_year_type(all_rows),
        "painel_empresas_rais": painel_empresas_rais,
        "metodologia": pd.DataFrame(
            [
                {
                    "tema": "I1",
                    "nota": "Numero de operacoes aprovadas contado como contratos por ano de aprovacao. A data de aprovacao vem da Base de QR FUST quando ha cruzamento por numero de operacao; quando nao ha data.aprovacao disponivel no QR, a planilha marca fallback pela data_da_contratacao na aba cobertura_data_aprovacao.",
                },
                {
                    "tema": "I2",
                    "nota": "Clientes apoiados contados por CPF/CNPJ normalizado. CNPJs mascarados e registros no formato mascara | CNPJ completo foram harmonizados.",
                },
                {
                    "tema": "I3",
                    "nota": "Pela regra do PDF, novos clientes exigem verificar ausencia de operacao aprovada em qualquer instrumento BNDES nos 36 meses anteriores. Como a base disponivel e FUST, a planilha sinaliza este valor como proxy pela primeira aparicao na Base de Operacoes_FUST_compilada.",
                },
                {
                    "tema": "I6",
                    "nota": "Variacao percentual calculada com estoque medio mensal anual da RAIS: (media anual do ano de referencia - media anual do ano anterior) / media anual do ano anterior. A tabela usa o painel completo das empresas apoiadas mapeadas na RAIS, independente do ano de apoio.",
                },
                {
                    "tema": "I7",
                    "nota": "Controle formado pelas empresas nao apoiadas selecionadas no matching e acompanhadas no painel RAIS anual. A variacao tambem usa estoque medio mensal anual reconstruido dos microdados RAIS_VINC.",
                },
                {
                    "tema": "Limite RAIS",
                    "nota": "A RAIS mensal reprocessada cobre 2022 a 2024. Assim, as variacoes anuais ficam disponiveis para 2023 e 2024. Nao ha RAIS 2025/2026 no disco para preencher I6/I7 desses anos sem extrapolacao.",
                },
                {
                    "tema": "I4, I5, I8 e I9",
                    "nota": "Nao observados diretamente nas bases usadas. Permanecem como linhas de disponibilidade para sinalizar lacuna de dados.",
                },
            ]
        ),
        "cobertura_rais": pd.DataFrame(
            [
                {
                    "empresas_base_operacoes": base["cnpj_key"].nunique(),
                    **rais_info,
                }
            ]
        ),
    }

    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        for sheet_name, data in sheets.items():
            clean = data.copy()
            clean.to_excel(writer, sheet_name=sheet_name[:31], index=False)

    style_workbook(xlsx_path)
    return xlsx_path


def style_workbook(path: Path) -> None:
    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="7F7F7F")
    header_font = Font(color="FFFFFF", bold=True)
    stripe_fill = PatternFill("solid", fgColor="F4F7F6")
    side = Side(style="thin", color="B7B7B7")
    border = Border(left=side, right=side, top=side, bottom=side)

    for ws in wb.worksheets:
        if ws.max_row < 1 or ws.max_column < 1:
            continue
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ws.max_column):
            if row[0].row % 2 == 0:
                for cell in row:
                    cell.fill = stripe_fill
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical="top", wrap_text=True)

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        headers = [str(ws.cell(row=1, column=col).value).lower() for col in range(1, ws.max_column + 1)]
        for col_idx, header in enumerate(headers, start=1):
            letter = get_column_letter(col_idx)
            if "valor_" in header:
                for cell in ws[letter][1:]:
                    cell.number_format = "#,##0.00"
            max_len = 10
            for cell in ws[letter][: min(ws.max_row, 100)]:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 55)

    wb.save(path)


def main() -> None:
    base = parse_base_operacoes()
    profile = company_profile(base)
    efficacy = efficacy_rows(base)
    years = sorted(base["ano_aprovacao"].dropna().astype(int).unique())
    tipos = [*sorted(base["tipo_operacao"].dropna().unique()), "Total"]
    matched_panel, rais_info = parse_painel_matching_rais()
    rais_panel, painel_empresas_rais = full_rais_panel(profile, matched_panel)
    effectiveness, match_audit = effectiveness_rows_full_panel(rais_panel)
    effectiveness = effectiveness[effectiveness["ano_aprovacao"].isin(years)].copy()
    match_audit = match_audit[match_audit["ano_aprovacao"].isin(years)].copy()
    effectiveness = add_missing_effectiveness_years(effectiveness, years, tipos)
    effectiveness = pd.concat(
        [effectiveness, unavailable_effectiveness_rows(years, tipos)],
        ignore_index=True,
        sort=False,
    )
    all_rows = pd.concat([efficacy, effectiveness], ignore_index=True, sort=False)
    summary = summarize_by_year(all_rows)
    path = save_outputs(all_rows, summary, match_audit, painel_empresas_rais, base, rais_info)
    print(path)
    print(f"linhas_por_ano={len(all_rows)}")
    print(f"linhas_matching={len(match_audit)}")


if __name__ == "__main__":
    main()
